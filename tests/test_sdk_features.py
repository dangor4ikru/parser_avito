"""Функции, перенесённые из eminsk/avito-sdk: характеристики, имя продавца, снижение цены."""
import gc
import tempfile
from pathlib import Path
from unittest.mock import MagicMock

from openpyxl import Workbook, load_workbook

from db_service import SQLiteDBHandler
from dto import AvitoConfig
from filters.ads_filter import AdsFilter
from integrations.notifications.vk import VKNotifier
from integrations.notifications.telegram import TelegramNotifier
from models import Item, PriceDetailed
from parser.export.excel import ExcelStorage
from parser.extractors import extract_params, extract_seller_name, format_params
from parser_cls import AvitoParse


def _price(value: int) -> PriceDetailed:
    return PriceDetailed(
        enabled=True, fullString=f"{value} ₽", hasValue=True, postfix="", string=f"{value} ₽",
        stringWithoutDiscount=None, title={}, titleDative="", value=value, wasLowered=False,
        exponent="",
    )


def _parser(db_handler, config=None):
    parser = AvitoParse.__new__(AvitoParse)
    parser.config = config or AvitoConfig(urls=["https://www.avito.ru/test"])
    parser.stop_event = None
    parser.good_request_count = 0
    parser.bad_request_count = 0
    parser.db_handler = db_handler
    return parser


# --- характеристики ---

def test_extract_params_beduin_scenario():
    payload = {"success": {"view": {"scenario": {"beduin": {"main": {"params": {
        "aboutRoom": {"items": [
            {"title": "Общая площадь", "description": "45.5 м²"},
            {"title": "Этаж", "description": "4 из 10"},
            {"title": "Пусто", "description": ""},
        ]},
    }}}}}}}
    assert extract_params(payload) == {"Общая площадь": "45.5 м²", "Этаж": "4 из 10"}


def test_extract_params_mobile_list_and_groups():
    payload = {"success": {"mobile": {
        "params": [{"title": "Состояние", "description": "Новое"}],
        "parameters": {
            "flat": [{"name": "Бренд", "value": "Apple"}],
            "groups": [{"items": [{"label": "Память", "text": "256 ГБ"}]}],
        },
    }}}
    assert extract_params(payload) == {"Состояние": "Новое", "Бренд": "Apple", "Память": "256 ГБ"}


def test_extract_params_garbage():
    assert extract_params(None) == {}
    assert extract_params({"success": {"mobile": {"params": "oops"}}}) == {}
    assert format_params(None) == ""
    assert format_params({"a": "1", "b": "2"}) == "a: 1; b: 2"


# --- имя продавца ---

def test_extract_seller_name_sources():
    assert extract_seller_name({"success": {"mobile": {"seller": {"name": " Алексей "}}}}) == "Алексей"
    assert extract_seller_name({"sellerName": "ComputerShop"}) == "ComputerShop"
    assert extract_seller_name({"sellerInfo": {"title": "re:Store"}}) == "re:Store"
    assert extract_seller_name({"sellerName": None, "sellerInfo": "?"}) is None


def test_catalog_seller_name_does_not_break_validation():
    ad = Item(id=1, sellerName={"unexpected": "shape"}, sellerInfo={"name": "Магазин"})
    parser = _parser(MagicMock())
    parser._add_seller_to_ads([ad])
    assert ad.seller_name == "Магазин"


def test_card_enrichment_fills_params_and_seller():
    config = AvitoConfig(urls=["https://www.avito.ru/test"], parse_params=True)
    parser = _parser(MagicMock(), config)
    parser.http = MagicMock()
    parser.http.fetch_item_data.return_value = {"success": {"mobile": {
        "seller": {"name": "Иван"},
        "params": [{"title": "Пробег", "description": "120 000 км"}],
    }}}
    ad = Item(id=5, urlPath="/item_5")
    parser.parse_params([ad])
    assert ad.characteristics == {"Пробег": "120 000 км"}
    assert ad.seller_name == "Иван"


def test_parse_params_disabled_makes_no_requests():
    parser = _parser(MagicMock())
    parser.http = MagicMock()
    parser.parse_params([Item(id=1)])
    parser.http.fetch_item_data.assert_not_called()


# --- снижение цены ---

def test_price_tracking_detects_drop_and_survives_restart():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp_dir:
        SQLiteDBHandler._instance = None
        db = SQLiteDBHandler(db_name=str(Path(tmp_dir) / "prices.db"))
        parser = _parser(db)

        first = Item(id=10, priceDetailed=_price(60000))
        parser._mark_price_changes([first])
        assert first.old_price is None
        parser._save_prices([first])

        same = Item(id=10, priceDetailed=_price(60000))
        parser._mark_price_changes([same])
        assert same.old_price is None

        cheaper = Item(id=10, priceDetailed=_price(52000))
        parser._mark_price_changes([cheaper])
        assert cheaper.old_price == 60000
        assert cheaper.price_drop == 8000
        parser._save_prices([cheaper])

        pricier = Item(id=10, priceDetailed=_price(55000))
        parser._mark_price_changes([pricier])
        assert pricier.old_price == 52000
        assert pricier.price_drop is None
        SQLiteDBHandler._instance = None
        gc.collect()


def test_price_tracking_falls_back_to_viewed_table():
    """База от старой версии: цены есть только в viewed."""
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp_dir:
        SQLiteDBHandler._instance = None
        db = SQLiteDBHandler(db_name=str(Path(tmp_dir) / "old.db"))
        db.add_record_from_page([Item(id=7, priceDetailed=_price(1000))])
        db.add_record_from_page([Item(id=7, priceDetailed=_price(900))])
        assert db.get_last_prices([7]) == {7: 900}
        SQLiteDBHandler._instance = None
        gc.collect()


def test_only_price_drops_filter():
    config = AvitoConfig(urls=["x"], only_price_drops=True, max_age=0, ignore_reserv=False)
    new_ad = Item(id=1, priceDetailed=_price(100))
    dropped = Item(id=2, priceDetailed=_price(80), old_price=100)
    raised = Item(id=3, priceDetailed=_price(120), old_price=100)
    assert AdsFilter(config).apply([new_ad, dropped, raised]) == [dropped]


# --- вывод ---

def test_notifications_show_seller_name_and_drop():
    ad = Item(id=2, title="PS5", priceDetailed=_price(52000), old_price=60000, seller_name="GameStore")
    tg = TelegramNotifier(bot_token="t", chat_id="c").format(ad)
    vk = VKNotifier.format_ad(ad)
    assert "GameStore" in tg and "8000" in tg and "60000" in tg
    assert "GameStore" in vk and "Цена снижена на 8000 (было 60000)" in vk


def test_notification_without_seller_has_no_none():
    ad = Item(id=3, title="Стол", priceDetailed=_price(500))
    assert "None" not in TelegramNotifier(bot_token="t", chat_id="c").format(ad)


def test_excel_new_columns_and_old_file_upgrade(tmp_path):
    path = tmp_path / "old.xlsx"
    wb = Workbook()
    wb.active.append(ExcelStorage.headers[:14])  # шапка старой версии
    wb.save(path)

    ad = Item(
        id=9, title="Квартира", priceDetailed=_price(5_000_000), old_price=5_500_000,
        seller_name="Агентство", characteristics={"Этаж": "4 из 10"},
    )
    ExcelStorage(path).save([ad])

    sheet = load_workbook(path).active
    header = [c.value for c in sheet[1]]
    row = dict(zip(header, [c.value for c in sheet[2]]))
    assert header == ExcelStorage.headers
    assert row["Имя продавца"] == "Агентство"
    assert row["Старая цена"] == 5_500_000
    assert row["Снижение цены"] == 500_000
    assert row["Характеристики"] == "Этаж: 4 из 10"
